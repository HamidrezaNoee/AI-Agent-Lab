# -*- coding: utf-8 -*-
"""Import ONE explicitly selected, FICTIONAL test document into AI.Document/DocumentChunk.

Supported: .txt / .md (UTF-8), .docx (requires python-docx).
This script MUST NOT be used for confidential or real company documents.
It writes only to EnterpriseProcess2_AI_Test with Windows Authentication.
"""
import argparse
import hashlib
import os
from pathlib import Path

import pyodbc

TEST_DB = "EnterpriseProcess2_AI_Test"
PREFIX = "SIM-MTG-A-%"
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_TEXT_CHARS = 250_000


def extract_text(path: Path) -> str:
    if path.suffix.lower() in (".txt", ".md"):
        text = path.read_text(encoding="utf-8-sig")
    elif path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise SystemExit("For DOCX files run: python -m pip install python-docx") from exc
        document = Document(str(path))
        pieces = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                pieces.append(" | ".join(cell.text for cell in row.cells))
        text = "\n\n".join(pieces)
    else:
        raise ValueError("Only .txt, .md and .docx are supported.")
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        raise ValueError("Document is empty after text extraction.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Extracted document is too long for this test importer.")
    return text


def chunk_text(text: str, size: int = 900, overlap: int = 140):
    if not (0 < overlap < size):
        raise ValueError("Invalid chunk size or overlap.")
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + size // 2, end),
                           text.rfind(" ", start + size // 2, end))
            if boundary > start + size // 2:
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            yield start, end, piece
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def main():
    parser = argparse.ArgumentParser(description="Ingest a SYNTHETIC test document.")
    parser.add_argument("--file", required=True, help="One explicit local file")
    parser.add_argument("--title", default="", help="Human-readable document title")
    parser.add_argument("--meeting-no", default="", help="Optional SIM-MTG-A-... number")
    args = parser.parse_args()
    path = Path(args.file).expanduser().resolve(strict=True)
    if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Not a regular file or exceeds 2 MB test-only limit.")
    text = extract_text(path)
    if len(path.name) > 260:
        raise ValueError("Filename exceeds schema limit.")
    title = args.title.strip() or path.stem
    if len(title) > 300:
        raise ValueError("Title exceeds schema limit.")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    chunks = list(chunk_text(text))
    server = os.getenv("SQL_SERVER", "")
    if not server:
        raise RuntimeError("Set SQL_SERVER to your test SQL Server instance.")
    if os.getenv("SQL_DATABASE") != TEST_DB:
        raise RuntimeError("SAFETY: SQL_DATABASE must be EnterpriseProcess2_AI_Test.")

    conn_str = (
        "DRIVER={ODBC Driver 17 for SQL Server};"
        f"SERVER={server};DATABASE={TEST_DB};"
        "Trusted_Connection=yes;Encrypt=yes;TrustServerCertificate=yes;"
    )
    conn = pyodbc.connect(conn_str, timeout=8, autocommit=False)
    try:
        cur = conn.cursor()
        cur.execute("SELECT DB_NAME()")
        if cur.fetchone()[0] != TEST_DB:
            raise RuntimeError("Wrong database. Ingestion cancelled.")
        meeting_id = None
        if args.meeting_no:
            if not args.meeting_no.startswith("SIM-MTG-A-"):
                raise ValueError("Only synthetic meeting numbers may be linked.")
            cur.execute(
                "SELECT MeetingID FROM MTG.Meeting WHERE MeetingNo=? AND MeetingNo LIKE ?",
                args.meeting_no, PREFIX,
            )
            meeting = cur.fetchone()
            if not meeting:
                raise ValueError("Synthetic meeting number not found.")
            meeting_id = int(meeting[0])

        cur.execute(
            """
            SELECT TOP (1) DocumentID, Sha256, VersionNo, RelatedMeetingID
            FROM AI.Document WITH (UPDLOCK, HOLDLOCK)
            WHERE OriginalFileName=?
            ORDER BY VersionNo DESC
            """,
            path.name,
        )
        previous = cur.fetchone()
        if previous and previous[1] == digest and previous[3] == meeting_id:
            conn.rollback()
            print(f"ALREADY IMPORTED: DocumentID={previous[0]}, Version={previous[2]}")
            return

        version = int(previous[2]) + 1 if previous else 1
        cur.execute(
            "UPDATE AI.Document SET IsCurrent=0 WHERE OriginalFileName=? AND IsCurrent=1",
            path.name,
        )
        cur.execute(
            """
            INSERT INTO AI.Document
                (OriginalFileName,Title,FileType,Sha256,VersionNo,
                 IsCurrent,IsSynthetic,RelatedMeetingID,CharCount)
            OUTPUT INSERTED.DocumentID
            VALUES (?,?,?,?,?,1,1,?,?)
            """,
            path.name, title, path.suffix.lower().lstrip("."),
            digest, version, meeting_id, len(text),
        )
        document_id = int(cur.fetchone()[0])
        for no, (start, end, piece) in enumerate(chunks, 1):
            cur.execute(
                """
                INSERT INTO AI.DocumentChunk
                    (DocumentID,ChunkNo,StartChar,EndChar,ChunkText)
                VALUES (?,?,?,?,?)
                """,
                document_id, no, start, end, piece,
            )
        conn.commit()
        print(f"IMPORTED SYNTHETIC TEST DOCUMENT: DocumentID={document_id}, "
              f"Version={version}, Chunks={len(chunks)}, RelatedMeetingID={meeting_id}")
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()
